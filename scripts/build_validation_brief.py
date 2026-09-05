"""Build MARKA's third-party validation brief as DOCX and PDF."""

from __future__ import annotations

import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    Image,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "validation" / "MARKA_third_party_validation_brief.md"
DOCX_OUTPUT = SOURCE.with_suffix(".docx")
PDF_OUTPUT = SOURCE.with_suffix(".pdf")
LOGO = ROOT / "assets" / "logos" / "marka_mark_100.png"

PLUM = "3B0042"
NAVY = "18243B"
MUTED = "697386"
WARM = "FBF7F2"
GOLD = "D9A441"
WHITE = "FFFFFF"


def source_body() -> list[str]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    first_rule = lines.index("---")
    return lines[first_rule + 1 :]


def markdown_inline_docx(paragraph, value: str) -> None:
    """Add a small Markdown subset (bold and inline code) to a DOCX paragraph."""
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", value)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(PLUM)
        else:
            paragraph.add_run(token)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("MARKA  •  Validation Brief     ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(19)
    section.right_margin = Mm(19)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Heading 1", 20, PLUM, 16, 7),
        ("Heading 2", 15, PLUM, 14, 6),
        ("Heading 3", 11, NAVY, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Pull Quote" not in [style.name for style in styles]:
        quote = styles.add_style("Pull Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["Pull Quote"]
    quote.font.name = "Liberation Sans"
    quote.font.size = Pt(12)
    quote.font.bold = True
    quote.font.color.rgb = RGBColor.from_string(PLUM)
    quote.paragraph_format.left_indent = Mm(7)
    quote.paragraph_format.right_indent = Mm(7)
    quote.paragraph_format.space_before = Pt(8)
    quote.paragraph_format.space_after = Pt(9)

    add_page_number(section.footer.paragraphs[0])


def add_docx_cover(document: Document) -> None:
    if LOGO.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Mm(18))
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run("MARKA")
    run.bold = True
    run.font.name = "Liberation Sans"
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor.from_string(PLUM)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("THINK ON PAPER. PERFORM UNDER PRESSURE.")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(GOLD)

    p = document.add_paragraph()
    run = p.add_run("Third-Party\nConcept Validation Brief")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Graduate aptitude preparation  |  5 September 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    box = document.add_table(rows=1, cols=1)
    box.autofit = False
    cell = box.cell(0, 0)
    shade_cell(cell, WARM)
    set_cell_margins(cell, 190, 220, 190, 220)
    p = cell.paragraphs[0]
    p.style = document.styles["Pull Quote"]
    p.paragraph_format.left_indent = Mm(0)
    p.paragraph_format.right_indent = Mm(0)
    p.add_run(
        "MARKA is a hybrid self-evaluation platform that teaches repeatable "
        "paper-working methods for aptitude problems, then tests whether those "
        "methods produce faster, accurate performance in timed CBT simulations."
    )

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(
        "For discussion with graduate applicants, career coaches, training providers, "
        "university career centres, and assessment professionals."
    )
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = document.add_paragraph()
    run = p.add_run("DISCUSSION DRAFT")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(PLUM)
    p.add_run("  •  Concept for validation—not a claim of proven outcomes")
    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build_docx() -> None:
    document = Document()
    configure_docx(document)
    document.core_properties.title = "MARKA Third-Party Concept Validation Brief"
    document.core_properties.subject = "Graduate aptitude preparation concept validation"
    document.core_properties.author = "MARKA"
    document.core_properties.comments = "Discussion draft; hypotheses require third-party validation."
    add_docx_cover(document)

    lines = source_body()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text:
            i += 1
            continue
        if text == "---":
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "5")
            bottom.set(qn("w:color"), "E4DDE5")
            borders.append(bottom)
            pPr.append(borders)
            i += 1
            continue
        if text.startswith("## "):
            if text.startswith("## 14."):
                document.add_page_break()
            document.add_heading(text[3:], level=1)
            i += 1
            continue
        if text.startswith("### "):
            document.add_heading(text[4:], level=2)
            i += 1
            continue
        if text.startswith("| "):
            rows, i = parse_table(lines, i)
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = True
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if row_index == 0:
                        shade_cell(cell, PLUM)
                    p = cell.paragraphs[0]
                    markdown_inline_docx(p, value)
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor.from_string(WHITE)
                    if row_index > 0 and row_index % 2 == 0:
                        shade_cell(cell, "F6F2F6")
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if text.startswith("> "):
            p = document.add_paragraph(style="Pull Quote")
            markdown_inline_docx(p, text[2:])
            i += 1
            continue
        if re.match(r"^- ", text):
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Mm(6)
            p.paragraph_format.first_line_indent = Mm(-3)
            markdown_inline_docx(p, text[2:])
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", text)
        if numbered:
            p = document.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Mm(6)
            p.paragraph_format.first_line_indent = Mm(-3)
            markdown_inline_docx(p, numbered.group(2))
            i += 1
            continue

        paragraph_lines = [text]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate or candidate == "---" or candidate.startswith(("## ", "### ", "| ", "> ", "- ")):
                break
            if re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        p = document.add_paragraph()
        markdown_inline_docx(p, " ".join(paragraph_lines))

    document.save(DOCX_OUTPUT)


def register_pdf_fonts() -> tuple[str, str, str]:
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    regular = font_dir / "DejaVuSans.ttf"
    bold = font_dir / "DejaVuSans-Bold.ttf"
    italic = font_dir / "DejaVuSans-Oblique.ttf"
    if regular.exists() and bold.exists() and italic.exists():
        pdfmetrics.registerFont(TTFont("MarkaSans", str(regular)))
        pdfmetrics.registerFont(TTFont("MarkaSans-Bold", str(bold)))
        pdfmetrics.registerFont(TTFont("MarkaSans-Italic", str(italic)))
        pdfmetrics.registerFontFamily(
            "MarkaSans",
            normal="MarkaSans",
            bold="MarkaSans-Bold",
            italic="MarkaSans-Italic",
            boldItalic="MarkaSans-Bold",
        )
        return "MarkaSans", "MarkaSans-Bold", "MarkaSans-Italic"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def pdf_inline(value: str) -> str:
    escaped = escape(value)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`(.+?)`", r'<font name="Courier" color="#3B0042">\1</font>', escaped)
    return escaped


def build_pdf() -> None:
    regular, bold, italic = register_pdf_fonts()
    stylesheet = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle(
            "Body",
            parent=stylesheet["BodyText"],
            fontName=regular,
            fontSize=8.7,
            leading=12.1,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=4.5,
        ),
        "h1": ParagraphStyle(
            "H1",
            parent=stylesheet["Heading1"],
            fontName=bold,
            fontSize=16.5,
            leading=19.5,
            textColor=colors.HexColor(f"#{PLUM}"),
            spaceBefore=11,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "H2",
            parent=stylesheet["Heading2"],
            fontName=bold,
            fontSize=11.3,
            leading=14,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "quote": ParagraphStyle(
            "Quote",
            parent=stylesheet["BodyText"],
            fontName=bold,
            fontSize=11,
            leading=15,
            leftIndent=7 * mm,
            rightIndent=7 * mm,
            textColor=colors.HexColor(f"#{PLUM}"),
            spaceBefore=7,
            spaceAfter=8,
        ),
        "small": ParagraphStyle(
            "Small",
            parent=stylesheet["BodyText"],
            fontName=regular,
            fontSize=7.3,
            leading=10,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
        "cover_title": ParagraphStyle(
            "CoverTitle",
            parent=stylesheet["Title"],
            fontName=bold,
            fontSize=29,
            leading=33,
            alignment=TA_LEFT,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=10,
        ),
    }

    doc = BaseDocTemplate(
        str(PDF_OUTPUT),
        pagesize=A4,
        rightMargin=19 * mm,
        leftMargin=19 * mm,
        topMargin=18 * mm,
        bottomMargin=17 * mm,
        title="MARKA Third-Party Concept Validation Brief",
        author="MARKA",
        subject="Graduate aptitude preparation concept validation",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")

    def footer(canvas, current_doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#E4DDE5"))
        canvas.setLineWidth(0.4)
        canvas.line(19 * mm, 13.5 * mm, A4[0] - 19 * mm, 13.5 * mm)
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(19 * mm, 9 * mm, "MARKA  •  Third-Party Concept Validation Brief")
        canvas.drawRightString(A4[0] - 19 * mm, 9 * mm, str(current_doc.page))
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="content", frames=[frame], onPage=footer)])
    story = []

    if LOGO.exists():
        story.append(Image(str(LOGO), width=18 * mm, height=18 * mm))
    story.append(Spacer(1, 19 * mm))
    story.append(Paragraph("MARKA", ParagraphStyle(
        "Wordmark", fontName=bold, fontSize=34, leading=38,
        textColor=colors.HexColor(f"#{PLUM}"), spaceAfter=3,
    )))
    story.append(Paragraph(
        "THINK ON PAPER. PERFORM UNDER PRESSURE.",
        ParagraphStyle("Tagline", fontName=bold, fontSize=9.5, leading=12,
                       textColor=colors.HexColor(f"#{GOLD}"), spaceAfter=17),
    ))
    story.append(Paragraph("Third-Party<br/>Concept Validation Brief", styles["cover_title"]))
    story.append(Paragraph(
        "Graduate aptitude preparation &nbsp;&nbsp;|&nbsp;&nbsp; 5 September 2026",
        ParagraphStyle("CoverMeta", fontName=regular, fontSize=9, leading=12,
                       textColor=colors.HexColor(f"#{MUTED}"), spaceAfter=15),
    ))
    callout = Table([[Paragraph(
        "MARKA is a hybrid self-evaluation platform that teaches repeatable "
        "paper-working methods for aptitude problems, then tests whether those "
        "methods produce faster, accurate performance in timed CBT simulations.",
        styles["quote"],
    )]], colWidths=[doc.width])
    callout.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{WARM}")),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#E4DDE5")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4 * mm),
    ]))
    story.extend([callout, Spacer(1, 11 * mm)])
    story.append(Paragraph(
        "For discussion with graduate applicants, career coaches, training providers, "
        "university career centres, and assessment professionals.",
        ParagraphStyle("CoverAudience", fontName=regular, fontSize=10, leading=14,
                       textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=8),
    ))
    story.append(Paragraph(
        "<b>DISCUSSION DRAFT</b> &nbsp;•&nbsp; Concept for validation—not a claim of proven outcomes",
        styles["small"],
    ))
    story.append(PageBreak())

    lines = source_body()
    i = 0
    bullets: list[ListItem] = []
    bullet_kind = None

    def flush_bullets():
        nonlocal bullets, bullet_kind
        if bullets:
            story.append(ListFlowable(
                bullets,
                bulletType="bullet" if bullet_kind == "bullet" else "1",
                start="1",
                leftIndent=13,
                bulletFontName=regular,
                bulletFontSize=7,
                bulletColor=colors.HexColor(f"#{PLUM}"),
                spaceAfter=4,
            ))
        bullets = []
        bullet_kind = None

    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text:
            flush_bullets()
            i += 1
            continue
        if text == "---":
            flush_bullets()
            story.append(HRFlowable(width="100%", thickness=0.45, color=colors.HexColor("#E4DDE5"),
                                    spaceBefore=3, spaceAfter=3))
            i += 1
            continue
        if text.startswith("## "):
            flush_bullets()
            if text.startswith("## 14."):
                story.append(PageBreak())
            story.append(Paragraph(pdf_inline(text[3:]), styles["h1"]))
            i += 1
            continue
        if text.startswith("### "):
            flush_bullets()
            story.append(Paragraph(pdf_inline(text[4:]), styles["h2"]))
            i += 1
            continue
        if text.startswith("| "):
            flush_bullets()
            rows, i = parse_table(lines, i)
            data = [[Paragraph(pdf_inline(cell), styles["small"]) for cell in row] for row in rows]
            if data:
                if len(data[0]) == 2:
                    widths = [doc.width * 0.55, doc.width * 0.45]
                else:
                    widths = [doc.width / len(data[0])] * len(data[0])
                table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
                commands = [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{PLUM}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), bold),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDD5DE")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
                for row_index in range(2, len(data), 2):
                    commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F7F3F7")))
                table.setStyle(TableStyle(commands))
                story.extend([table, Spacer(1, 4)])
            continue
        if text.startswith("> "):
            flush_bullets()
            story.append(Paragraph(pdf_inline(text[2:]), styles["quote"]))
            i += 1
            continue
        if text.startswith("- "):
            if bullet_kind not in (None, "bullet"):
                flush_bullets()
            bullet_kind = "bullet"
            bullets.append(ListItem(Paragraph(pdf_inline(text[2:]), styles["body"]), leftIndent=8))
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", text)
        if numbered:
            if bullet_kind not in (None, "number"):
                flush_bullets()
            bullet_kind = "number"
            bullets.append(ListItem(Paragraph(pdf_inline(numbered.group(2)), styles["body"]), leftIndent=8))
            i += 1
            continue

        flush_bullets()
        paragraph_lines = [text]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate or candidate == "---" or candidate.startswith(("## ", "### ", "| ", "> ", "- ")):
                break
            if re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        story.append(Paragraph(pdf_inline(" ".join(paragraph_lines)), styles["body"]))

    flush_bullets()
    doc.build(story)


def main() -> None:
    build_docx()
    build_pdf()
    print(DOCX_OUTPUT)
    print(PDF_OUTPUT)


if __name__ == "__main__":
    main()
